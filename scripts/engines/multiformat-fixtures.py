"""Development fixtures for the opt-in MultiFormatRealTest; run with private Python.

XLS generation uses xlwt installed only in build/multiformat/fixture-tools.
The Outlook fixture is from the upstream Microsoft MarkItDown test suite.
"""
from pathlib import Path
import sys
import subprocess
import urllib.request
import zipfile

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / 'build/multiformat/fixtures'
OUT.mkdir(parents=True, exist_ok=True)
sys.path.insert(0, str(ROOT / 'build/multiformat/fixture-tools'))
from openpyxl import Workbook
from pptx import Presentation
import xlwt
from PIL import Image, ImageDraw, ImageFont

def target(extension):
    return OUT / ('documento ñ.' + extension)

subprocess.run([sys.executable, str(ROOT / 'scripts/spikes/new-marker-fixture.py'), str(target('pdf'))], check=True)
with zipfile.ZipFile(target('docx'), 'w') as z:
    z.writestr('[Content_Types].xml', '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Default Extension="xml" ContentType="application/xml"/><Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/></Types>')
    z.writestr('_rels/.rels', '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/></Relationships>')
    z.writestr('word/document.xml', '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body><w:p><w:pPr><w:pStyle w:val="Heading1"/></w:pPr><w:r><w:t>ParseForge Word</w:t></w:r></w:p><w:p><w:r><w:t>Contenido de prueba</w:t></w:r></w:p></w:body></w:document>')
with zipfile.ZipFile(target('epub'), 'w') as z:
    z.writestr('mimetype', 'application/epub+zip')
    z.writestr('META-INF/container.xml', '<container xmlns="urn:oasis:names:tc:opendocument:xmlns:container"><rootfiles><rootfile full-path="OEBPS/content.opf" media-type="application/oebps-package+xml"/></rootfiles></container>')
    z.writestr('OEBPS/content.opf', '<package xmlns="http://www.idpf.org/2007/opf" version="2.0" unique-identifier="id"><metadata xmlns:dc="http://purl.org/dc/elements/1.1/"><dc:title>ParseForge EPUB</dc:title><dc:identifier id="id">test</dc:identifier><dc:language>es</dc:language></metadata><manifest><item id="ch1" href="chapter.xhtml" media-type="application/xhtml+xml"/></manifest><spine><itemref idref="ch1"/></spine></package>')
    z.writestr('OEBPS/chapter.xhtml', '<html xmlns="http://www.w3.org/1999/xhtml"><body><h1>ParseForge EPUB</h1><p>Contenido de prueba</p></body></html>')
deck = Presentation()
slide = deck.slides.add_slide(deck.slide_layouts[1])
slide.shapes.title.text = 'ParseForge PowerPoint'
slide.placeholders[1].text = 'Contenido de prueba'
deck.save(target('pptx'))
book = Workbook()
book.active.append(['ParseForge', 'Valor'])
book.active.append(['Ejemplo', 42])
book.save(target('xlsx'))
old = xlwt.Workbook()
sheet = old.add_sheet('ParseForge')
sheet.write(0, 0, 'ParseForge XLS')
sheet.write(1, 0, 42)
old.save(str(target('xls')))
for extension, content in {
    'html': '<h1>ParseForge HTML</h1><ul><li>Contenido</li></ul>',
    'htm': '<h1>ParseForge HTM</h1>', 'txt': 'ParseForge texto', 'md': '# ParseForge Markdown',
    'csv': 'Nombre,Valor\nParseForge,42\n', 'json': '{"nombre":"ParseForge","valor":42}',
    'xml': '<documento><nombre>ParseForge</nombre></documento>'}.items():
    target(extension).write_text(content, encoding='utf-8')
with zipfile.ZipFile(target('zip'), 'w') as z:
    z.write(target('docx'), 'documento.docx')
    z.write(target('txt'), 'documento.txt')
urllib.request.urlretrieve('https://raw.githubusercontent.com/microsoft/markitdown/main/packages/markitdown/tests/test_files/test_outlook_msg.msg', target('msg'))
image = Image.new('RGB', (1000, 300), 'white')
ImageDraw.Draw(image).text((50, 80), 'ParseForge image OCR', fill='black', font=ImageFont.truetype('C:/Windows/Fonts/arial.ttf', 48))
image.save(target('png'))
print(OUT)
